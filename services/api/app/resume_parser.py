"""Bounded offline extraction. No OCR, skill inference, networking or automatic scoring."""

import io
import json
import os

# The child uses a fixed isolated interpreter, with document bytes on stdin only.
import subprocess  # nosec B404
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

PARSER_VERSION = "text-v1-pypdf6.21.2-docx1.2.0"
MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 100_000
MAX_PAGES = 20
ERRORS = {
    "unsupported_type",
    "size_limit",
    "invalid_document",
    "encrypted_pdf",
    "empty_text",
    "resource_limit",
}


@dataclass(frozen=True)
class ParseResult:
    status: str
    text: str | None = None
    error_code: str | None = None
    parser_version: str = PARSER_VERSION


class Rejected(Exception):
    pass


def extract_resume(data: bytes, extension: str) -> ParseResult:
    """Never treat a parsing failure as a zero-score or an empty confirmed resume."""
    if extension not in {"pdf", "docx"}:
        return ParseResult("unscored", error_code="unsupported_type")
    if not data or len(data) > MAX_BYTES:
        return ParseResult("unscored", error_code="size_limit")
    # -I ignores PYTHON* variables/user site packages. No cloud credentials reach the child.
    environment = {key: os.environ[key] for key in ("SYSTEMROOT", "WINDIR") if key in os.environ}
    try:
        with tempfile.TemporaryDirectory(prefix="smarthire-parser-") as directory:
            # The executable and script are fixed; the extension is an allowlisted enum.
            process = subprocess.run(  # nosec B603
                [sys.executable, "-I", str(Path(__file__).resolve()), extension],
                input=data,
                stdout=subprocess.PIPE,
                stderr=subprocess.DEVNULL,
                timeout=15,
                check=False,
                cwd=directory,
                env=environment,
            )
        if process.returncode != 0 or len(process.stdout) > 800_000:
            return ParseResult("unscored", error_code="resource_limit")
        result = json.loads(process.stdout)
        text, code = result.get("text"), result.get("error_code")
        if code in ERRORS:
            return ParseResult("unscored", error_code=code)
        if (
            code is not None
            or not isinstance(text, str)
            or not text.strip()
            or len(text) > MAX_TEXT
        ):
            return ParseResult("unscored", error_code="invalid_document")
        return ParseResult("parsed", text=text)
    except subprocess.TimeoutExpired:
        return ParseResult("unscored", error_code="resource_limit")
    except Exception:  # noqa: BLE001 - do not leak parser diagnostics or document text.
        return ParseResult("unscored", error_code="invalid_document")


def _extract_pdf(data):
    from pypdf import PdfReader

    if not data.startswith(b"%PDF-"):
        raise Rejected("invalid_document")
    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise Rejected("encrypted_pdf")
    if not 1 <= len(reader.pages) <= MAX_PAGES:
        raise Rejected("size_limit")
    parts = []
    size = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        size += len(text)
        if size > MAX_TEXT:
            raise Rejected("size_limit")
        parts.append(text)
    return "\n".join(parts)


def _extract_docx(data):
    from docx import Document
    from lxml import etree

    if not data.startswith(b"PK\x03\x04"):
        raise Rejected("invalid_document")
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        members = package.infolist()
        names = {member.filename for member in members}
        if (
            len(members) > 200
            or len(names) != len(members)
            or not {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= names
        ):
            raise Rejected("invalid_document")
        if sum(member.file_size for member in members) > 20 * 1024 * 1024:
            raise Rejected("size_limit")
        for member in members:
            name = PurePosixPath(member.filename)
            if (
                name.is_absolute()
                or ".." in name.parts
                or "\\" in member.filename
                or "vbaproject" in member.filename.lower()
                or member.flag_bits & 1
            ):
                raise Rejected("invalid_document")
            if (
                member.file_size > 4 * 1024 * 1024
                or member.file_size > max(member.compress_size, 1) * 100
            ):
                raise Rejected("size_limit")
            if member.filename.endswith((".xml", ".rels")):
                parser = etree.XMLParser(
                    resolve_entities=False, load_dtd=False, no_network=True, huge_tree=False
                )
                tree = etree.fromstring(package.read(member), parser=parser).getroottree()
                if tree.docinfo.doctype:
                    raise Rejected("invalid_document")
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(parts)


def _child_main():
    try:
        if os.name == "posix":
            import resource

            resource.setrlimit(resource.RLIMIT_CPU, (8, 8))
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        data = sys.stdin.buffer.read(MAX_BYTES + 1)
        if not data or len(data) > MAX_BYTES:
            raise Rejected("size_limit")
        kind = sys.argv[1]
        if kind not in {"pdf", "docx"}:
            raise Rejected("unsupported_type")
        text = _extract_pdf(data) if kind == "pdf" else _extract_docx(data)
        text = unicodedata.normalize("NFKC", text).replace("\x00", "").strip()
        if len(text) > MAX_TEXT:
            raise Rejected("size_limit")
        if not text:
            raise Rejected("empty_text")
        print(json.dumps({"text": text}))
    except Rejected as error:
        code = str(error)
        print(json.dumps({"error_code": code if code in ERRORS else "invalid_document"}))
    except Exception:  # noqa: BLE001 - document and library diagnostics never escape.
        print(json.dumps({"error_code": "invalid_document"}))


if __name__ == "__main__":
    _child_main()
