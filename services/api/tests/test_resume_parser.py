import io
import json
import subprocess
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.resume_parser import MAX_BYTES, PARSER_VERSION, extract_resume


def pdf_bytes(text=None, pages=1, encrypted=False):
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=400, height=400)
        if text is not None:
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {
                            NameObject("/F1"): DictionaryObject(
                                {
                                    NameObject("/Type"): NameObject("/Font"),
                                    NameObject("/Subtype"): NameObject("/Type1"),
                                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                                }
                            )
                        }
                    )
                }
            )
            content = DecodedStreamObject()
            content.set_data(b"BT /F1 12 Tf 20 20 Td (" + text.encode() + b") Tj ET")
            page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("nonpersonal-fixture")
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def docx_bytes():
    document = Document()
    document.add_paragraph("Nonpersonal Flutter developer fixture")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Dart"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def altered_docx(name, content):
    target = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(docx_bytes())) as source,
        zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive,
    ):
        for member in source.infolist():
            if member.filename != name:
                archive.writestr(member.filename, source.read(member))
        archive.writestr(name, content)
    return target.getvalue()


def test_pdf_text_is_extracted_offline_with_version_and_no_score():
    result = extract_resume(pdf_bytes("Nonpersonal Flutter Python fixture"), "pdf")
    assert result.status == "parsed" and "Flutter Python" in result.text
    assert result.parser_version == PARSER_VERSION and result.error_code is None
    assert not hasattr(result, "score")


def test_docx_paragraphs_and_tables_are_extracted():
    result = extract_resume(docx_bytes(), "docx")
    assert result.status == "parsed"
    assert "Flutter developer" in result.text and "Python\tDart" in result.text


@pytest.mark.parametrize(
    "data,kind,code",
    [
        (b"not-a-resume", "exe", "unsupported_type"),
        (b"", "pdf", "size_limit"),
        (b"x" * (MAX_BYTES + 1), "pdf", "size_limit"),
        (b"not a PDF", "pdf", "invalid_document"),
        (b"not a ZIP", "docx", "invalid_document"),
    ],
    ids=["unsupported", "empty", "oversize", "wrong-pdf-signature", "wrong-docx-signature"],
)
def test_rejected_input_is_explicitly_unscored(data, kind, code):
    result = extract_resume(data, kind)
    assert result.status == "unscored" and result.text is None and result.error_code == code


def test_encrypted_and_image_only_pdfs_are_not_scored():
    assert extract_resume(pdf_bytes("secret", encrypted=True), "pdf").error_code == "encrypted_pdf"
    assert extract_resume(pdf_bytes(), "pdf").error_code == "empty_text"


def test_pdf_page_limit_applies_before_extraction():
    assert extract_resume(pdf_bytes(pages=21), "pdf").error_code == "size_limit"


@pytest.mark.parametrize(
    "name,content,code",
    [
        ("../escape.xml", b"<root/>", "invalid_document"),
        ("word/vbaProject.bin", b"macro", "invalid_document"),
        ("word/large.xml", b"x" * (2 * 1024 * 1024), "size_limit"),
        (
            "word/document.xml",
            b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///nonpersonal-never-open">]><x>&e;</x>',
            "invalid_document",
        ),
        (
            "word/document.xml",
            '<!DOCTYPE x [<!ENTITY e SYSTEM "https://example.invalid/never-fetch">]><x>&e;</x>'.encode(
                "utf-16"
            ),
            "invalid_document",
        ),
    ],
    ids=["traversal", "macro", "zip-bomb", "external-file-entity", "utf16-external-entity"],
)
def test_docx_archive_and_entity_boundaries(name, content, code):
    result = extract_resume(altered_docx(name, content), "docx")
    assert result.status == "unscored" and result.text is None and result.error_code == code


def test_child_has_no_cloud_credentials_or_user_environment(monkeypatch):
    observed = {}
    monkeypatch.setenv("GOOGLE_APPLICATION_CREDENTIALS", "private-never-forwarded")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "private-never-forwarded")

    def child(command, **options):
        observed.update(options)
        assert command[1] == "-I" and command[-1] == "pdf"
        return SimpleNamespace(
            returncode=0, stdout=json.dumps({"text": "Nonpersonal text"}).encode()
        )

    monkeypatch.setattr("app.resume_parser.subprocess.run", child)
    assert extract_resume(b"fixture", "pdf").status == "parsed"
    assert "GOOGLE_APPLICATION_CREDENTIALS" not in observed["env"]
    assert "AWS_SECRET_ACCESS_KEY" not in observed["env"]
    assert observed["stderr"] == subprocess.DEVNULL and observed["timeout"] == 15


def test_timeout_and_malformed_child_output_fail_closed(monkeypatch):
    def timeout(*args, **options):
        raise subprocess.TimeoutExpired("fixed-parser", 15)

    monkeypatch.setattr("app.resume_parser.subprocess.run", timeout)
    assert extract_resume(b"fixture", "pdf").error_code == "resource_limit"
    monkeypatch.setattr(
        "app.resume_parser.subprocess.run",
        lambda *a, **k: SimpleNamespace(returncode=0, stdout=b"private malformed diagnostics"),
    )
    result = extract_resume(b"fixture", "pdf")
    assert result.status == "unscored" and result.text is None
